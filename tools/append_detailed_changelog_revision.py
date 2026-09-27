"""Create a protected, append-only DOCX revision from a Markdown addendum.

Example:
  python tools/append_detailed_changelog_revision.py docs/append/Rev005-Cobertura.md

The tool never edits an earlier revision. It verifies that every paragraph from
the previous DOCX remains an identical prefix, adds a new appendix, enables
Word's read-only protection, then asks the hash-chain registrar to append the
new revision record.
"""
from datetime import date
from pathlib import Path
from copy import deepcopy
import os
import re
import stat
import subprocess
import sys
import tempfile
import zipfile
from lxml import etree

from docx import Document
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
REV_RE = re.compile(r"CalradiaForge-Registro-Mejoras-Rev(\d{3})\.docx$")
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
CANONICAL_SETTINGS_NS = {
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "o": "urn:schemas-microsoft-com:office:office",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "v": "urn:schemas-microsoft-com:vml",
    "w10": "urn:schemas-microsoft-com:office:word",
    "w": W_NS,
    "w14": "http://schemas.microsoft.com/office/word/2010/wordml",
    "sl": "http://schemas.openxmlformats.org/schemaLibrary/2006/main",
}
SETTINGS_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.settings+xml"


def read_docx_paragraphs(path):
    return [(paragraph.text, paragraph.style.name) for paragraph in Document(path).paragraphs]


def read_body_paragraph_xml(path):
    with zipfile.ZipFile(path) as package:
        root = etree.fromstring(package.read("word/document.xml"))
    body = root.find(f"{{{W_NS}}}body")
    if body is None:
        raise ValueError(f"No se encontró w:body en {path.name}")
    return [etree.tostring(child, method="c14n") for child in body if child.tag == f"{{{W_NS}}}p"]


def _settings_namespace_map(original_settings):
    """Keep every source prefix because mc:Ignorable stores prefixes as text."""
    nsmap = {
        prefix: uri
        for prefix, uri in original_settings.nsmap.items()
        if prefix != "xml"
    }
    for preferred_prefix, uri in CANONICAL_SETTINGS_NS.items():
        if preferred_prefix not in nsmap or nsmap[preferred_prefix] == uri:
            nsmap[preferred_prefix] = uri
            continue
        if uri in nsmap.values():
            continue
        alias = f"cf_{preferred_prefix}"
        suffix = 2
        while alias in nsmap:
            alias = f"cf_{preferred_prefix}{suffix}"
            suffix += 1
        nsmap[alias] = uri
    return nsmap


def set_read_only_protection(docx_path):
    """Set the same non-password Word read-only protection used by the document skill."""
    temp_path = docx_path.with_name(docx_path.name + ".protecting")
    with zipfile.ZipFile(docx_path, "r") as source:
        entries = [(info, source.read(info.filename)) for info in source.infolist()]
    replacements = {}
    names = {info.filename for info, _ in entries}
    settings_data = next((data for info, data in entries if info.filename == "word/settings.xml"), None)
    if settings_data is None:
        settings = etree.Element(f"{{{W_NS}}}settings", nsmap=CANONICAL_SETTINGS_NS)
    else:
        original_settings = etree.fromstring(settings_data)
        # Word's mc:Ignorable value contains literal prefixes. Preserve source
        # aliases verbatim (even when two prefixes share a URI); add canonical
        # aliases without replacing a source binding that happens to use the
        # same prefix for a different namespace.
        nsmap = _settings_namespace_map(original_settings)
        ignorable = original_settings.get(f"{{{CANONICAL_SETTINGS_NS['mc']}}}Ignorable", "").split()
        unresolved = [prefix for prefix in ignorable if prefix not in nsmap]
        if unresolved:
            raise ValueError(f"No se pueden resolver prefijos mc:Ignorable en settings.xml: {unresolved}")
        settings = etree.Element(original_settings.tag, nsmap=nsmap)
        for name, value in original_settings.attrib.items():
            settings.set(name, value)
        for child in original_settings:
            settings.append(deepcopy(child))
    for node in list(settings.iter(f"{{{W_NS}}}documentProtection")):
        node.getparent().remove(node)
    protection = etree.Element(f"{{{W_NS}}}documentProtection")
    protection.set(f"{{{W_NS}}}edit", "readOnly")
    protection.set(f"{{{W_NS}}}enforcement", "1")
    protection.set(f"{{{W_NS}}}formatting", "0")
    settings.insert(0, protection)
    replacements["word/settings.xml"] = etree.tostring(settings, encoding="UTF-8", xml_declaration=True, standalone=True)

    with zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as target:
        for info, data in entries:
            target.writestr(info, replacements.get(info.filename, data))
        for name, data in replacements.items():
            if name not in names:
                target.writestr(name, data)
    os.replace(temp_path, docx_path)
    docx_path.chmod(stat.S_IREAD)


def append_markdown(doc, markdown_path, revision):
    lines = markdown_path.read_text(encoding="utf-8-sig").splitlines()
    if not any(line.strip() for line in lines):
        raise ValueError("El anexo Markdown está vacío")
    doc.add_page_break()
    title = f"Anexo Rev{revision:03} — agregado {date.today().strftime('%d-%m-%Y')}"
    doc.add_heading(title, level=1)
    intro = doc.add_paragraph()
    run = intro.add_run("Regla de anexado. ")
    run.bold = True
    run.font.color.rgb = RGBColor(120, 91, 31)
    intro.add_run("Este anexo amplía el registro anterior; no reemplaza ni corrige sus párrafos. Las precisiones históricas se atribuyen a los changelogs fuente.")

    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=2)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=3)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            doc.add_paragraph(line)


def main():
    if len(sys.argv) != 2:
        raise SystemExit("Uso: python append_detailed_changelog_revision.py <anexo.md>")
    annex = Path(sys.argv[1]).expanduser().resolve()
    if not annex.is_file():
        raise FileNotFoundError(f"No existe el anexo: {annex}")
    lock_path = DOCS / ".CalradiaForge-Registro-Mejoras.append.lock"
    try:
        lock_fd = os.open(lock_path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
    except FileExistsError as exc:
        raise RuntimeError(f"Ya hay un anexado activo o quedó un lock que requiere revisión manual: {lock_path}") from exc
    try:
        os.write(lock_fd, f"pid={os.getpid()}\n".encode("ascii"))
        os.close(lock_fd)
        revisions = {}
        for path in DOCS.glob("CalradiaForge-Registro-Mejoras-Rev*.docx"):
            match = REV_RE.fullmatch(path.name)
            if match:
                revisions[int(match.group(1))] = path
        if not revisions:
            raise RuntimeError("No existe la revisión base protegida")
        if sorted(revisions) != list(range(1, max(revisions) + 1)):
            raise RuntimeError("Hay huecos entre las revisiones existentes")
        previous_number = max(revisions)
        previous_path = revisions[previous_number]
        revision = previous_number + 1
        output = DOCS / f"CalradiaForge-Registro-Mejoras-Rev{revision:03}.docx"
        if output.exists():
            raise FileExistsError(f"No se sobrescribe ninguna revisión: {output}")
        before = read_docx_paragraphs(previous_path)
        before_xml = read_body_paragraph_xml(previous_path)
        doc = Document(previous_path)
        append_markdown(doc, annex, revision)
        with tempfile.NamedTemporaryFile(prefix=f"CalradiaForge-Rev{revision:03}-", suffix=".docx", dir=DOCS, delete=False) as stream:
            pending = Path(stream.name)
        try:
            doc.save(pending)
            after = read_docx_paragraphs(pending)
            if after[:len(before)] != before:
                raise RuntimeError("La comprobación de prefijo detectó cambios en el texto o estilo anterior")
            after_xml = read_body_paragraph_xml(pending)
            if after_xml[:len(before_xml)] != before_xml:
                raise RuntimeError("La comprobación OOXML detectó cambios en párrafos anteriores")
            set_read_only_protection(pending)
            os.rename(pending, output)
        finally:
            if pending.exists():
                pending.chmod(stat.S_IWRITE | stat.S_IREAD)
                pending.unlink()

        registrar = ROOT / "tools" / "record_detailed_changelog_integrity.py"
        subprocess.run([sys.executable, str(registrar)], cwd=ROOT, check=True)
        print(f"Creada revisión append-only Rev{revision:03}: {output}")
        print(f"Párrafos previos preservados literalmente: {len(before)}; anexo: {annex}")
    finally:
        if lock_path.exists():
            lock_path.unlink()


if __name__ == "__main__":
    main()
